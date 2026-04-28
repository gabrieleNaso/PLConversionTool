from __future__ import annotations

import json
import re
from pathlib import Path

try:
    from openpyxl import load_workbook
except ModuleNotFoundError:  # pragma: no cover - runtime guard
    load_workbook = None


TEMPLATE_RELATIVE_PATH = Path("docs/templates/ir_excel_template_single_page_with_support_fc.xlsx")


def _project_root() -> Path:
    # src/plc_converter/excel_export.py -> repo root
    return Path(__file__).resolve().parents[2]


def _template_path() -> Path:
    return (_project_root() / TEMPLATE_RELATIVE_PATH).resolve()


def _infer_step_number(step_name: str) -> int | None:
    match = re.match(r"^S(\d+)$", step_name.strip(), flags=re.IGNORECASE)
    if not match:
        return None
    value = int(match.group(1))
    return value if value > 0 else None


def _collect_expression_operands(expression: str) -> list[str]:
    if not expression:
        return []
    reserved = {"AND", "OR", "NOT", "TRUE", "FALSE"}
    found: list[str] = []
    for token in re.findall(r"[A-Za-z_]\\w*(?:\\.\\w+)*", expression):
        if token.upper() in reserved or token in found:
            continue
        found.append(token)
    return found


def _normalize_support_category(category: str) -> str:
    value = (category or "").strip().lower()
    aliases = {
        "lv2": "mode",
        "lev2": "mode",
        "transition": "transitions",
        "transizioni": "transitions",
    }
    return aliases.get(value, value)


def _category_from_memory_role(role: str) -> str:
    raw = (role or "").strip().lower()
    if raw in {"hmi"}:
        return "hmi"
    if raw in {"external"}:
        return "external"
    if raw in {"mode", "lv2", "lev2"}:
        return "lv2"
    if raw in {"transitions", "transition"}:
        return "transitions"
    if raw in {"alarm", "fault"}:
        return "alarm"
    if raw in {"output"}:
        return "output"
    return "aux"


def _control_kind_from_awl_timer(kind: str) -> str:
    raw = (kind or "").strip().upper()
    if raw in {"SD", "TON", "T_ON"}:
        return "t_on"
    if raw in {"SF", "TOF", "T_OFF"}:
        return "t_off"
    if raw in {"SE", "SP", "SS", "TP", "T_P"}:
        return "t_p"
    return "t_on"


def export_excel_from_ir(ir_payload: dict, output_path: Path) -> None:
    """
    Export an IR JSON payload to an Excel workbook compatible with the manual IR template.

    The layout is intentionally stable because it is used as part of the project pipeline
    (AWL -> IR JSON -> Excel -> XML generation).
    """

    if load_workbook is None:
        raise RuntimeError(
            "Modulo mancante: openpyxl. Installa le dipendenze backend (pip install -r backend/requirements.txt)."
        )

    template_path = _template_path()
    if not template_path.exists():
        raise RuntimeError(f"Template Excel non trovato: {template_path}")

    wb = load_workbook(template_path)
    required_sheets = {"sequence", "operands", "support_fc"}
    missing = sorted(required_sheets - set(wb.sheetnames))
    if missing:
        raise RuntimeError(
            "Template Excel non valido: mancano fogli obbligatori: " + ", ".join(missing)
        )

    # Ensure output workbook matches template format: keep only required sheets.
    for sheet_name in list(wb.sheetnames):
        if sheet_name not in required_sheets:
            del wb[sheet_name]

    ws_sequence = wb["sequence"]
    ws_operands = wb["operands"]
    ws_support = wb["support_fc"]
    ws_passthrough = wb.create_sheet("ir_passthrough")
    ws_passthrough.sheet_state = "hidden"
    ws_passthrough.append(["json_chunk"])

    # Clear any pre-filled example rows, but keep headers and formatting.
    for ws in (ws_sequence, ws_operands, ws_support):
        if ws.max_row and ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)

    def header_map(ws) -> dict[str, int]:
        values = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        mapping: dict[str, int] = {}
        for idx, value in enumerate(values, start=1):
            key = str(value).strip() if value is not None else ""
            if key:
                mapping[key] = idx
        return mapping

    seq_cols = header_map(ws_sequence)
    op_cols = header_map(ws_operands)
    sup_cols = header_map(ws_support)

    def append_dict_row(ws, cols: dict[str, int], values: dict[str, object]) -> None:
        ws.append([None] * max(cols.values(), default=0))
        row_idx = ws.max_row
        for key, col_idx in cols.items():
            if key in values:
                ws.cell(row=row_idx, column=col_idx, value=values.get(key))

    steps = ir_payload.get("steps", []) if isinstance(ir_payload.get("steps"), list) else []
    transitions = (
        ir_payload.get("transitions", [])
        if isinstance(ir_payload.get("transitions"), list)
        else []
    )
    step_numbers: dict[str, int | None] = {}
    for step in steps:
        name = str(step.get("name") or "").strip()
        if not name:
            continue
        raw_no = step.get("step_number")
        try:
            step_no = int(raw_no) if raw_no is not None else _infer_step_number(name)
        except (TypeError, ValueError):
            step_no = _infer_step_number(name)
        step_numbers[name] = step_no

    for idx, transition in enumerate(transitions, start=1):
        source = str(transition.get("source_step") or "").strip()
        target = str(transition.get("target_step") or "").strip()
        append_dict_row(
            ws_sequence,
            seq_cols,
            {
                "step_name": source or target,
                "numero_step": step_numbers.get(source)
                or _infer_step_number(source or target or "")
                or "",
                "from_step": source,
                "transition_id": str(transition.get("transition_id") or f"T{idx}"),
                "to_step": target,
                "condition_expression": str(transition.get("guard_expression") or "TRUE"),
                "flow_type": str(transition.get("flow_type") or "alternative"),
                "parallel_group": str(transition.get("parallel_group") or ""),
            },
        )

    # Steps detail is represented in 'sequence' sheet in the template format.

    operand_categories = (
        dict(ir_payload.get("operand_categories", {}))
        if isinstance(ir_payload.get("operand_categories"), dict)
        else {}
    )
    operand_datatypes = (
        dict(ir_payload.get("operand_datatypes", {}))
        if isinstance(ir_payload.get("operand_datatypes"), dict)
        else {}
    )
    operand_notes = (
        dict(ir_payload.get("operand_notes", {}))
        if isinstance(ir_payload.get("operand_notes"), dict)
        else {}
    )
    operand_controls = (
        dict(ir_payload.get("operand_control_settings", {}))
        if isinstance(ir_payload.get("operand_control_settings"), dict)
        else {}
    )

    ordered_operands: list[str] = []

    # Prefer DB-facing members so the operands sheet reflects what ends up materialized
    # in the generated support DBs (useful for manual Excel editing).
    derived_catalog_built = False
    try:
        from plc_converter.analysis import (  # type: ignore
            _collect_aux_support_members,
            _collect_diag_support_members,
            _collect_external_support_members,
            _collect_hmi_support_members,
            _collect_io_support_members,
            _collect_mode_support_members,
            _collect_output_family_members,
            _collect_parameters_support_members,
            _collect_timer_trigger_support_members_by_category,
            _collect_transitions_support_members,
            _ir_from_payload,
        )

        ir_obj = _ir_from_payload(ir_payload=ir_payload)
        derived_members_by_category: dict[str, list[str]] = {
            "alarm": [name for name, _ in _collect_diag_support_members(ir_obj)],
            "hmi": [name for name, _ in _collect_hmi_support_members(ir_obj)],
            "aux": [name for name, _ in _collect_aux_support_members(ir_obj)],
            "transitions": [
                name for name, _ in _collect_transitions_support_members(ir_obj, [])
            ],
            "output": [name for name, _ in _collect_output_family_members(ir_obj)]
            + [name for name, _ in _collect_io_support_members(ir_obj)],
            "lv2": [name for name, _ in _collect_mode_support_members(ir_obj)],
            "external": [name for name, _ in _collect_external_support_members(ir_obj)],
        }
        # Parameters DB lives under AUX family but doesn't have a dedicated Excel category.
        derived_members_by_category["aux"].extend(
            [name for name, _ in _collect_parameters_support_members(ir_obj)]
        )
        trigger_members = _collect_timer_trigger_support_members_by_category(ir_obj)
        trigger_category_map = {
            "diag": "alarm",
            "hmi": "hmi",
            "aux": "aux",
            "transitions": "transitions",
            "io": "output",
        }
        for raw_category, members in trigger_members.items():
            excel_category = trigger_category_map.get(raw_category, "aux")
            derived_members_by_category.setdefault(excel_category, []).extend(
                [name for name, _ in members]
            )

        for excel_category, members in derived_members_by_category.items():
            for name in members:
                token = str(name or "").strip()
                if not token:
                    continue
                if token not in ordered_operands:
                    ordered_operands.append(token)
                operand_categories.setdefault(token, excel_category)
        derived_catalog_built = True
    except Exception:
        pass

    # Append any explicit/raw operands that aren't already covered by the DB-facing catalog.
    explicit_catalog = ir_payload.get("operand_catalog", [])
    if isinstance(explicit_catalog, list):
        for operand in explicit_catalog:
            name = str(operand).strip()
            if name and name not in ordered_operands:
                ordered_operands.append(name)

    for transition in transitions:
        guard_operands = transition.get("guard_operands", [])
        if not isinstance(guard_operands, list):
            guard_operands = []
        for operand in guard_operands:
            name = str(operand).strip()
            if name and name not in ordered_operands:
                ordered_operands.append(name)
        for operand in _collect_expression_operands(
            str(transition.get("guard_expression") or "")
        ):
            if operand and operand not in ordered_operands:
                ordered_operands.append(operand)

    for fault in ir_payload.get("faults", []):
        if not isinstance(fault, dict):
            continue
        name = str(fault.get("name") or "").strip()
        if name and name not in ordered_operands:
            ordered_operands.append(name)
        if name:
            operand_categories.setdefault(name, "alarm")

    for out in ir_payload.get("outputs", []):
        if not isinstance(out, dict):
            continue
        name = str(out.get("name") or "").strip()
        if name and name not in ordered_operands:
            ordered_operands.append(name)
        if name:
            operand_categories.setdefault(name, "output")

    for mem in ir_payload.get("memories", []):
        if not isinstance(mem, dict):
            continue
        name = str(mem.get("name") or "").strip()
        if name and name not in ordered_operands:
            ordered_operands.append(name)
        if name:
            operand_categories.setdefault(
                name, _category_from_memory_role(str(mem.get("role") or ""))
            )

    # If we already materialized the DB-facing catalog, avoid leaking raw DB addresses
    # for external references into the operands sheet.
    if not derived_catalog_built:
        for ext in ir_payload.get("external_refs", []):
            name = str(ext).strip()
            if name and name not in ordered_operands:
                ordered_operands.append(name)
            if name:
                operand_categories.setdefault(name, "external")

    for timer in ir_payload.get("timers", []):
        if not isinstance(timer, dict):
            continue
        source = str(timer.get("source_timer") or "").strip()
        if source and source not in ordered_operands:
            ordered_operands.append(source)
        if source:
            operand_datatypes.setdefault(source, "IEC_TIMER")
            operand_controls.setdefault(
                source,
                {
                    "kind": _control_kind_from_awl_timer(str(timer.get("kind") or "")),
                    "value": str(timer.get("preset") or "T#1S").replace("S5T#", "T#"),
                },
            )
            operand_categories.setdefault(source, "aux")
        for trig in (
            timer.get("trigger_operands", [])
            if isinstance(timer.get("trigger_operands"), list)
            else []
        ):
            trig_name = str(trig).strip()
            if trig_name and trig_name not in ordered_operands:
                ordered_operands.append(trig_name)

    # Other IR sections are intentionally omitted from the template workbook.

    for operand in ordered_operands:
        category = str(operand_categories.get(operand) or "")
        datatype = str(operand_datatypes.get(operand) or "")
        note = str(operand_notes.get(operand) or "")
        control = operand_controls.get(operand) if isinstance(operand_controls, dict) else None
        if not isinstance(control, dict):
            control = {}
        append_dict_row(
            ws_operands,
            op_cols,
            {
                "operand": operand,
                "category": category,
                "datatype": datatype,
                "control_kind": str(control.get("kind") or ""),
                "control_value": str(control.get("value") or ""),
                "note": note,
            },
        )

    wrote_support = False
    saw_excel_support_rows = False
    for member in ir_payload.get("support_members", []):
        if not isinstance(member, dict):
            continue
        category = _normalize_support_category(str(member.get("category") or ""))
        member_name = str(member.get("member_name") or "")
        if not category or not member_name:
            continue
        saw_excel_support_rows = True
        append_dict_row(
            ws_support,
            sup_cols,
            {
                "category": category,
                "member_name": member_name,
                "result_member": "",
                "condition_expression": "",
                "coil_mode": "",
                "comment": str(member.get("comment") or ""),
                "network": member.get("network_index") or "",
            },
        )
        wrote_support = True

    for logic in ir_payload.get("support_logic", []):
        if not isinstance(logic, dict):
            continue
        if "result_member" not in logic or "category" not in logic:
            continue
        category = _normalize_support_category(str(logic.get("category") or ""))
        result_member = str(logic.get("result_member") or "")
        if not category or not result_member:
            continue
        saw_excel_support_rows = True
        append_dict_row(
            ws_support,
            sup_cols,
            {
                "category": category,
                "member_name": "",
                "result_member": result_member,
                "condition_expression": str(logic.get("condition_expression") or "TRUE"),
                "coil_mode": str(logic.get("coil_mode") or ""),
                "comment": str(logic.get("comment") or ""),
                "network": logic.get("network_index") or "",
            },
        )
        wrote_support = True

    # When IR comes from AWL, support_members/support_logic may contain analysis metadata
    # (not Excel rows). In that case, derive the same member set that the generator would
    # normally infer so the visible sheet isn't empty.
    if not saw_excel_support_rows:
        try:
            from plc_converter.analysis import (  # type: ignore
                _collect_aux_support_members,
                _collect_diag_support_members,
                _collect_external_support_members,
                _collect_hmi_support_members,
                _collect_io_support_members,
                _collect_mode_support_members,
                _collect_output_family_members,
                _collect_parameters_support_members,
                _collect_timer_trigger_support_members_by_category,
                _collect_transitions_support_members,
                _derive_awl_action_logic_rows,
                _derive_awl_timer_logic_rows,
                _ir_from_payload,
            )

            ir_obj = _ir_from_payload(ir_payload=ir_payload)
            derived: dict[str, list[tuple[str, str]]] = {
                "diag": _collect_diag_support_members(ir_obj),
                "hmi": _collect_hmi_support_members(ir_obj),
                "aux": _collect_aux_support_members(ir_obj),
                "transitions": _collect_transitions_support_members(ir_obj, []),
                "io": _collect_io_support_members(ir_obj),
                "output": _collect_output_family_members(ir_obj),
                "mode": _collect_mode_support_members(ir_obj),
                "external": _collect_external_support_members(ir_obj),
                "parameters": _collect_parameters_support_members(ir_obj),
            }
            trigger_members = _collect_timer_trigger_support_members_by_category(ir_obj)
            for category, members in trigger_members.items():
                derived.setdefault(category, []).extend(members)

            seen_pairs: set[tuple[str, str]] = set()
            for category, members in derived.items():
                normalized_cat = _normalize_support_category(category)
                if not normalized_cat:
                    continue
                for member_name, comment in members:
                    key = (normalized_cat, str(member_name))
                    if not member_name or key in seen_pairs:
                        continue
                    seen_pairs.add(key)
                    append_dict_row(
                        ws_support,
                        sup_cols,
                        {
                            "category": normalized_cat,
                            "member_name": str(member_name),
                            "result_member": "",
                            "condition_expression": "",
                            "coil_mode": "",
                            "comment": str(comment or ""),
                            "network": "",
                        },
                    )
                    wrote_support = True

            # Also materialize derived logic rows (timers/actions) so the visible support_fc
            # resembles what the bundle will contain even when IR originates from AWL.
            derived_actions = _derive_awl_action_logic_rows(ir_obj)
            timer_logic = _derive_awl_timer_logic_rows(ir_obj)
            for row in timer_logic:
                append_dict_row(
                    ws_support,
                    sup_cols,
                    {
                        "category": "aux",
                        "member_name": "",
                        "result_member": str(row.get("result_member") or ""),
                        "condition_expression": str(row.get("condition_expression") or "TRUE"),
                        "coil_mode": str(row.get("coil_mode") or ""),
                        "comment": str(row.get("comment") or ""),
                        "network": row.get("network_index") or "",
                    },
                )
                wrote_support = True

            for category in ("aux", "io", "hmi", "diag", "transitions"):
                for row in derived_actions.get(category, []):
                    append_dict_row(
                        ws_support,
                        sup_cols,
                        {
                            "category": category,
                            "member_name": "",
                            "result_member": str(row.get("result_member") or ""),
                            "condition_expression": str(row.get("condition_expression") or "TRUE"),
                            "coil_mode": str(row.get("coil_mode") or ""),
                            "comment": str(row.get("comment") or ""),
                            "network": row.get("network_index") or "",
                        },
                    )
                    wrote_support = True
        except Exception:
            # Best-effort fallback: keep transitions-derived rows below.
            pass

    # Keep FC sheet populated even when support_logic contains only metadata.
    for idx, transition in enumerate(transitions, start=1):
        transition_id = str(transition.get("transition_id") or f"T{idx}").strip()
        guard_expression = (
            str(transition.get("guard_expression") or "TRUE").strip() or "TRUE"
        )
        network_index = transition.get("network_index") or idx
        append_dict_row(
            ws_support,
            sup_cols,
            {
                "category": "transitions",
                "member_name": "",
                "result_member": transition_id,
                "condition_expression": guard_expression,
                "coil_mode": "",
                "comment": f"Derived from IR transition {transition_id}",
                "network": network_index,
            },
        )
        wrote_support = True

    if not wrote_support:
        append_dict_row(
            ws_support,
            sup_cols,
            {
                "category": "transitions",
                "member_name": "",
                "result_member": "Transitions_Enable",
                "condition_expression": "TRUE",
                "coil_mode": "",
                "comment": "",
                "network": 1,
            },
        )

    # 1:1 roundtrip support (Excel -> IR JSON): keep the original payload embedded.
    # scripts/generate_from_excel_ir.py will prefer this sheet when present.
    passthrough_json = json.dumps(ir_payload, ensure_ascii=False)
    chunk_size = 30000
    for start in range(0, len(passthrough_json), chunk_size):
        ws_passthrough.append([passthrough_json[start : start + chunk_size]])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
