#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
except ModuleNotFoundError:  # pragma: no cover - runtime guard
    Workbook = None


PROJECT_ROOT = Path(__file__).resolve().parents[1]
BACKEND_ROOT = PROJECT_ROOT / "backend"
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from app.core_converter import analyze_conversion, export_conversion_bundle_from_ir  # noqa: E402
from app.core_converter import analyze_conversion_project  # noqa: E402


SUPPORTED_EXTENSIONS = {".awl", ".txt", ".md"}


def _infer_step_number(step_name: str) -> int | None:
    match = re.match(r"^S(\d+)$", step_name.strip(), flags=re.IGNORECASE)
    if not match:
        return None
    value = int(match.group(1))
    return value if value > 0 else None


def _collect_expression_operands(expression: str) -> list[str]:
    if not expression:
        return []
    reserved = {"AND", "OR", "NOT", "TRUE", "FALSE"}
    found: list[str] = []
    for token in re.findall(r"[A-Za-z_]\w*(?:\.\w+)*", expression):
        if token.upper() in reserved or token in found:
            continue
        found.append(token)
    return found


def _infer_category_from_operand(operand: str) -> str:
    upper = operand.upper()
    if upper.startswith("Q"):
        return "output"
    if upper.startswith("DB12"):
        return "hmi"
    if upper.startswith("DB11"):
        return "alarm"
    if upper.startswith("DB13"):
        return "aux"
    if upper.startswith("DB14"):
        return "transitions"
    if upper.startswith("DB17"):
        return "lv2"
    if upper.startswith("DB18"):
        return "external"
    if upper.startswith("M"):
        return "aux"
    return "aux"


def _normalize_support_category(category: str) -> str:
    value = (category or "").strip().lower()
    aliases = {
        "lv2": "mode",
        "lev2": "mode",
        "transition": "transitions",
        "transizioni": "transitions",
    }
    return aliases.get(value, value)


def _category_from_memory_role(role: str) -> str:
    raw = (role or "").strip().lower()
    if raw in {"hmi"}:
        return "hmi"
    if raw in {"external"}:
        return "external"
    if raw in {"mode", "lv2", "lev2"}:
        return "lv2"
    if raw in {"transitions", "transition"}:
        return "transitions"
    if raw in {"alarm", "fault"}:
        return "alarm"
    if raw in {"output"}:
        return "output"
    return "aux"


def _control_kind_from_awl_timer(kind: str) -> str:
    raw = (kind or "").strip().upper()
    if raw in {"SD", "TON", "T_ON"}:
        return "t_on"
    if raw in {"SF", "TOF", "T_OFF"}:
        return "t_off"
    if raw in {"SE", "SP", "SS", "TP", "T_P"}:
        return "t_p"
    return "t_on"


def _export_excel_from_ir(ir_payload: dict, output_path: Path) -> None:
    if Workbook is None:
        raise RuntimeError(
            "Modulo mancante: openpyxl. Installa le dipendenze backend (pip install -r backend/requirements.txt)."
        )

    wb = Workbook()
    ws_sequence = wb.active
    ws_sequence.title = "sequence"
    ws_operands = wb.create_sheet("operands")
    ws_support = wb.create_sheet("support_fc")
    ws_passthrough = wb.create_sheet("ir_passthrough")
    ws_steps = wb.create_sheet("ir_steps")
    ws_transitions = wb.create_sheet("ir_transitions")
    ws_timers = wb.create_sheet("ir_timers")
    ws_memories = wb.create_sheet("ir_memories")
    ws_faults = wb.create_sheet("ir_faults")
    ws_outputs = wb.create_sheet("ir_outputs")
    ws_external = wb.create_sheet("ir_external_refs")
    ws_support_members = wb.create_sheet("ir_support_members")
    ws_support_logic = wb.create_sheet("ir_support_logic")
    ws_summary = wb.create_sheet("ir_summary")

    ws_sequence.append(
        [
            "step_name",
            "numero_step",
            "transition_id",
            "from_step",
            "to_step",
            "condition_expression",
            "flow_type",
            "parallel_group",
        ]
    )
    ws_operands.append(["operand", "category", "datatype", "control_kind", "control_value", "note"])
    ws_support.append(
        [
            "category",
            "member_name",
            "result_member",
            "condition_expression",
            "comment",
            "network",
        ]
    )
    ws_passthrough.append(["json_chunk"])
    ws_steps.append(["name", "step_number", "source_networks", "activation_networks", "action_networks"])
    ws_transitions.append(
        [
            "transition_id",
            "source_step",
            "target_step",
            "network_index",
            "guard_expression",
            "guard_operands",
            "jump_labels",
            "flow_type",
            "parallel_group",
        ]
    )
    ws_timers.append(["source_timer", "network_index", "kind", "preset", "trigger_operands"])
    ws_memories.append(["name", "role", "network_index"])
    ws_faults.append(["name", "network_index", "evidence"])
    ws_outputs.append(["name", "network_index", "action"])
    ws_external.append(["name"])
    ws_support_members.append(["category", "member_name", "comment", "network_index", "network_title"])
    ws_support_logic.append(
        [
            "category",
            "network_index",
            "network_title",
            "result_member",
            "condition_expression",
            "condition_operands",
            "coil_mode",
            "comment",
            "kind",
            "raw_json",
        ]
    )
    ws_summary.append(["field", "value"])

    steps = ir_payload.get("steps", []) if isinstance(ir_payload.get("steps"), list) else []
    transitions = ir_payload.get("transitions", []) if isinstance(ir_payload.get("transitions"), list) else []
    step_numbers: dict[str, int | None] = {}
    for step in steps:
        name = str(step.get("name") or "").strip()
        if not name:
            continue
        raw_no = step.get("step_number")
        try:
            step_no = int(raw_no) if raw_no is not None else _infer_step_number(name)
        except (TypeError, ValueError):
            step_no = _infer_step_number(name)
        step_numbers[name] = step_no

    for idx, transition in enumerate(transitions, start=1):
        source = str(transition.get("source_step") or "").strip()
        target = str(transition.get("target_step") or "").strip()
        ws_sequence.append(
            [
                source or target,
                step_numbers.get(source) or _infer_step_number(source or target or "") or "",
                str(transition.get("transition_id") or f"T{idx}"),
                source,
                target,
                str(transition.get("guard_expression") or "TRUE"),
                str(transition.get("flow_type") or "alternative"),
                str(transition.get("parallel_group") or ""),
            ]
        )
        ws_transitions.append(
            [
                str(transition.get("transition_id") or f"T{idx}"),
                source,
                target,
                transition.get("network_index") or "",
                str(transition.get("guard_expression") or "TRUE"),
                " | ".join(str(x) for x in (transition.get("guard_operands") if isinstance(transition.get("guard_operands"), list) else [])),
                " | ".join(str(x) for x in (transition.get("jump_labels") if isinstance(transition.get("jump_labels"), list) else [])),
                str(transition.get("flow_type") or ""),
                str(transition.get("parallel_group") or ""),
            ]
        )

    for step in steps:
        ws_steps.append(
            [
                str(step.get("name") or ""),
                step.get("step_number") if step.get("step_number") is not None else "",
                " | ".join(str(x) for x in (step.get("source_networks") if isinstance(step.get("source_networks"), list) else [])),
                " | ".join(str(x) for x in (step.get("activation_networks") if isinstance(step.get("activation_networks"), list) else [])),
                " | ".join(str(x) for x in (step.get("action_networks") if isinstance(step.get("action_networks"), list) else [])),
            ]
        )

    operand_categories = dict(ir_payload.get("operand_categories", {})) if isinstance(ir_payload.get("operand_categories"), dict) else {}
    operand_datatypes = dict(ir_payload.get("operand_datatypes", {})) if isinstance(ir_payload.get("operand_datatypes"), dict) else {}
    operand_notes = dict(ir_payload.get("operand_notes", {})) if isinstance(ir_payload.get("operand_notes"), dict) else {}
    operand_controls = dict(ir_payload.get("operand_control_settings", {})) if isinstance(ir_payload.get("operand_control_settings"), dict) else {}

    ordered_operands: list[str] = []
    explicit_catalog = ir_payload.get("operand_catalog", [])
    if isinstance(explicit_catalog, list):
        for operand in explicit_catalog:
            name = str(operand).strip()
            if name and name not in ordered_operands:
                ordered_operands.append(name)

    for transition in transitions:
        guard_operands = transition.get("guard_operands", [])
        if not isinstance(guard_operands, list):
            guard_operands = []
        for operand in guard_operands:
            name = str(operand).strip()
            if name and name not in ordered_operands:
                ordered_operands.append(name)
        for operand in _collect_expression_operands(str(transition.get("guard_expression") or "")):
            if operand and operand not in ordered_operands:
                ordered_operands.append(operand)

    # Preserve non-transition symbols so roundtrip AWL->Excel->XML stays consistent.
    for fault in ir_payload.get("faults", []):
        if not isinstance(fault, dict):
            continue
        name = str(fault.get("name") or "").strip()
        if name and name not in ordered_operands:
            ordered_operands.append(name)
        if name:
            operand_categories.setdefault(name, "alarm")

    for out in ir_payload.get("outputs", []):
        if not isinstance(out, dict):
            continue
        name = str(out.get("name") or "").strip()
        if name and name not in ordered_operands:
            ordered_operands.append(name)
        if name:
            operand_categories.setdefault(name, "output")

    for mem in ir_payload.get("memories", []):
        if not isinstance(mem, dict):
            continue
        name = str(mem.get("name") or "").strip()
        if name and name not in ordered_operands:
            ordered_operands.append(name)
        if name:
            operand_categories.setdefault(name, _category_from_memory_role(str(mem.get("role") or "")))

    for ext in ir_payload.get("external_refs", []):
        name = str(ext).strip()
        if name and name not in ordered_operands:
            ordered_operands.append(name)
        if name:
            operand_categories.setdefault(name, "external")

    for timer in ir_payload.get("timers", []):
        if not isinstance(timer, dict):
            continue
        source = str(timer.get("source_timer") or "").strip()
        if source and source not in ordered_operands:
            ordered_operands.append(source)
        if source:
            operand_datatypes.setdefault(source, "IEC_TIMER")
            operand_controls.setdefault(
                source,
                {
                    "kind": _control_kind_from_awl_timer(str(timer.get("kind") or "")),
                    "value": str(timer.get("preset") or "T#1S").replace("S5T#", "T#"),
                },
            )
            operand_categories.setdefault(source, "aux")
        for trig in timer.get("trigger_operands", []) if isinstance(timer.get("trigger_operands"), list) else []:
            trig_name = str(trig).strip()
            if trig_name and trig_name not in ordered_operands:
                ordered_operands.append(trig_name)
        ws_timers.append(
            [
                source,
                timer.get("network_index") or "",
                str(timer.get("kind") or ""),
                str(timer.get("preset") or ""),
                " | ".join(str(x) for x in (timer.get("trigger_operands") if isinstance(timer.get("trigger_operands"), list) else [])),
            ]
        )

    for mem in ir_payload.get("memories", []):
        if not isinstance(mem, dict):
            continue
        ws_memories.append([str(mem.get("name") or ""), str(mem.get("role") or ""), mem.get("network_index") or ""])

    for fault in ir_payload.get("faults", []):
        if not isinstance(fault, dict):
            continue
        ws_faults.append([str(fault.get("name") or ""), fault.get("network_index") or "", str(fault.get("evidence") or "")])

    for out in ir_payload.get("outputs", []):
        if not isinstance(out, dict):
            continue
        ws_outputs.append([str(out.get("name") or ""), out.get("network_index") or "", str(out.get("action") or "")])

    for ref in ir_payload.get("external_refs", []):
        ws_external.append([str(ref)])

    if not ordered_operands:
        ordered_operands.append("M0.0")

    for operand in ordered_operands:
        control = operand_controls.get(operand, {}) if isinstance(operand_controls.get(operand), dict) else {}
        ws_operands.append(
            [
                operand,
                str(operand_categories.get(operand) or _infer_category_from_operand(operand)),
                str(operand_datatypes.get(operand) or "Bool"),
                str(control.get("kind") or ""),
                str(control.get("value") or ""),
                str(operand_notes.get(operand) or ""),
            ]
        )

    wrote_support = False
    for member in ir_payload.get("support_members", []):
        if not isinstance(member, dict):
            continue
        category = _normalize_support_category(str(member.get("category") or ""))
        member_name = str(member.get("member_name") or "")
        if not category or not member_name:
            continue
        ws_support.append(
            [
                category,
                member_name,
                "",
                "",
                str(member.get("comment") or ""),
                member.get("network_index") or "",
            ]
        )
        ws_support_members.append(
            [
                category,
                member_name,
                str(member.get("comment") or ""),
                member.get("network_index") or "",
                str(member.get("network_title") or ""),
            ]
        )
        wrote_support = True

    for logic in ir_payload.get("support_logic", []):
        if not isinstance(logic, dict):
            continue
        if "result_member" not in logic or "category" not in logic:
            continue
        category = _normalize_support_category(str(logic.get("category") or ""))
        result_member = str(logic.get("result_member") or "")
        if not category or not result_member:
            continue
        ws_support.append(
            [
                category,
                "",
                result_member,
                str(logic.get("condition_expression") or "TRUE"),
                str(logic.get("comment") or ""),
                logic.get("network_index") or "",
            ]
        )
        ws_support_logic.append(
            [
                category,
                logic.get("network_index") or "",
                str(logic.get("network_title") or ""),
                result_member,
                str(logic.get("condition_expression") or "TRUE"),
                " | ".join(str(x) for x in (logic.get("condition_operands") if isinstance(logic.get("condition_operands"), list) else [])),
                str(logic.get("coil_mode") or ""),
                str(logic.get("comment") or ""),
                str(logic.get("kind") or ""),
                json.dumps(logic, ensure_ascii=False),
            ]
        )
        wrote_support = True

    # 1:1 readable FC projection: always materialize transition logic rows.
    # This keeps the FC sheet populated even when support_logic contains only metadata.
    for idx, transition in enumerate(transitions, start=1):
        transition_id = str(transition.get("transition_id") or f"T{idx}").strip()
        guard_expression = str(transition.get("guard_expression") or "TRUE").strip() or "TRUE"
        network_index = transition.get("network_index") or idx
        ws_support.append(
            [
                "transitions",
                "",
                transition_id,
                guard_expression,
                f"Derived from IR transition {transition_id}",
                network_index,
            ]
        )
        wrote_support = True

    if not wrote_support:
        ws_support.append(["transitions", "", "Transitions_Enable", "TRUE", "", 1])

    passthrough_json = json.dumps(ir_payload, ensure_ascii=False)
    chunk_size = 30000
    for start in range(0, len(passthrough_json), chunk_size):
        ws_passthrough.append([passthrough_json[start : start + chunk_size]])
    ws_summary.append(["sequence_name", str(ir_payload.get("sequence_name") or "")])
    ws_summary.append(["source_name", str(ir_payload.get("source_name") or "")])
    ws_summary.append(["steps_count", str(len(steps))])
    ws_summary.append(["transitions_count", str(len(transitions))])
    ws_summary.append(["timers_count", str(len(ir_payload.get("timers", []) if isinstance(ir_payload.get("timers"), list) else []))])
    ws_summary.append(["memories_count", str(len(ir_payload.get("memories", []) if isinstance(ir_payload.get("memories"), list) else []))])
    ws_summary.append(["faults_count", str(len(ir_payload.get("faults", []) if isinstance(ir_payload.get("faults"), list) else []))])
    ws_summary.append(["outputs_count", str(len(ir_payload.get("outputs", []) if isinstance(ir_payload.get("outputs"), list) else []))])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def _slugify(name: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9]+", "_", name).strip("_")
    return cleaned or "Sequence"


def _extract_awl_from_markdown(raw_text: str) -> str:
    fenced_blocks = re.findall(
        r"```(?:awl|il|stl|text)?\s*\n(.*?)```",
        raw_text,
        flags=re.IGNORECASE | re.DOTALL,
    )
    if not fenced_blocks:
        return raw_text

    selected_blocks = [block.strip() for block in fenced_blocks if _looks_like_awl_block(block)]
    if not selected_blocks:
        selected_blocks = [block.strip() for block in fenced_blocks]

    normalized_blocks: list[str] = []
    for index, block in enumerate(selected_blocks, start=1):
        if not block:
            continue
        if re.search(r"^\s*NETWORK\b", block, flags=re.IGNORECASE | re.MULTILINE):
            normalized_blocks.append(block)
        else:
            normalized_blocks.append(f"NETWORK {index}\n{block}")

    if not normalized_blocks:
        return raw_text
    return "\n\n".join(normalized_blocks).strip() + "\n"


def _looks_like_awl_block(block: str) -> bool:
    heuristics = (
        r"\bNETWORK\b",
        r"^\s*(A|AN|O|ON|U|UN|=|S|R|L|T|SD|SE|SP|SS|SF|JC|JCN|JU)\b",
        r"\bS5T#",
        r"\bDB\d+\.",
        r"\b[QAEIM]\d+\.\d+\b",
    )
    return any(
        re.search(pattern, block, flags=re.IGNORECASE | re.MULTILINE)
        for pattern in heuristics
    )


def _load_awl_text(path: Path) -> str:
    raw = path.read_text(encoding="utf-8")
    if path.suffix.lower() == ".md":
        return _extract_awl_from_markdown(raw)
    return raw


def _extract_block_id(path: Path, raw_text: str) -> str | None:
    # Prefer explicit markdown header: "# FC102 : ..." etc.
    header_match = re.search(
        r"^\s*#\s*(FC|FB|OB)\s*0*(\d+)\b",
        raw_text,
        flags=re.IGNORECASE | re.MULTILINE,
    )
    if header_match:
        return f"{header_match.group(1).upper()}{int(header_match.group(2))}"
    # Siemens AWL exports often start with FUNCTION/FUNCTION_BLOCK/ORGANIZATION_BLOCK.
    # Example: "FUNCTION FC 32 : ..." or "FUNCTION_BLOCK FB 10".
    decl_match = re.search(
        r"^\s*(?:FUNCTION|FUNCTION_BLOCK|ORGANIZATION_BLOCK)\s+(FC|FB|OB)\s*0*(\d+)\b",
        raw_text,
        flags=re.IGNORECASE | re.MULTILINE,
    )
    if decl_match:
        return f"{decl_match.group(1).upper()}{int(decl_match.group(2))}"
    # Fallback to filename if it embeds the block id.
    name_match = re.search(r"\b(FC|FB|OB)\s*0*(\d+)\b", path.stem, flags=re.IGNORECASE)
    if name_match:
        return f"{name_match.group(1).upper()}{int(name_match.group(2))}"
    return None


def _collect_sources(input_dir: Path, source: str | None, prefix: str | None) -> list[Path]:
    candidates = [
        item
        for item in input_dir.iterdir()
        if item.is_file() and item.suffix.lower() in SUPPORTED_EXTENSIONS
    ]
    if source:
        candidates = [item for item in candidates if item.name.lower() == source.lower()]
    if prefix:
        candidates = [item for item in candidates if item.name.lower().startswith(prefix.lower())]
    return sorted(candidates, key=lambda item: item.name.lower())


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Generate XML bundles from AWL files in an input folder."
    )
    parser.add_argument(
        "--input-dir",
        default="data/input",
        help="Input folder containing .awl/.txt/.md sources (default: data/input).",
    )
    parser.add_argument(
        "--output-root",
        default="data/output/generated",
        help="Output root folder (default: data/output/generated).",
    )
    parser.add_argument(
        "--name-prefix",
        default="Auto",
        help="Sequence name prefix used for generated bundles (default: Auto).",
    )
    parser.add_argument(
        "--source",
        default=None,
        help="Generate only this exact source filename from input dir.",
    )
    parser.add_argument(
        "--prefix",
        default=None,
        help="Generate only sources whose filename starts with this prefix.",
    )
    args = parser.parse_args()

    input_dir = (PROJECT_ROOT / args.input_dir).resolve()
    output_root = (PROJECT_ROOT / args.output_root).resolve()
    output_root.mkdir(parents=True, exist_ok=True)
    input_dir.mkdir(parents=True, exist_ok=True)

    sources = _collect_sources(input_dir, args.source, args.prefix)
    if not sources:
        print(f"No source files found in: {input_dir}")
        print("Add .awl, .txt, or .md files and run again.")
        return 0

    # Project scan: index all blocks available in the input folder so the converter
    # can resolve CALL dependencies (e.g. sequencer FC32) when those sources exist.
    project_blocks: dict[str, str] = {}
    for candidate in _collect_sources(input_dir, None, args.prefix):
        raw_text = candidate.read_text(encoding="utf-8")
        block_id = _extract_block_id(candidate, raw_text)
        if not block_id:
            continue
        project_blocks.setdefault(block_id, _extract_awl_from_markdown(raw_text) if candidate.suffix.lower() == ".md" else raw_text)

    generated = 0
    for source in sources:
        raw_text = source.read_text(encoding="utf-8")
        awl_source = _extract_awl_from_markdown(raw_text) if source.suffix.lower() == ".md" else raw_text
        if "NETWORK" not in awl_source.upper() and source.suffix.lower() != ".md":
            print(f"Skipping {source.name}: no AWL NETWORK found.")
            continue
        entry_block_id = _extract_block_id(source, raw_text)

        base_name = _slugify(source.stem)
        sequence_name = _slugify(f"{args.name_prefix}_{base_name}")
        bundle_dir = output_root / sequence_name.lower()
        if bundle_dir.exists():
            shutil.rmtree(bundle_dir)
        bundle_dir.mkdir(parents=True, exist_ok=True)
        bundle_dir_relative = bundle_dir.relative_to(PROJECT_ROOT)

        analysis = (
            analyze_conversion_project(
                sequence_name=sequence_name,
                awl_source=awl_source,
                source_name=source.name,
                project_blocks=project_blocks,
                entry_block_id=entry_block_id,
            )
            if project_blocks
            else analyze_conversion(
                sequence_name=sequence_name,
                awl_source=awl_source,
                source_name=source.name,
            )
        )
        ir_payload = analysis.get("ir")
        if not isinstance(ir_payload, dict):
            print(f"Skipping {source.name}: analysis IR non disponibile.")
            continue

        ir_json_path = bundle_dir / f"{sequence_name}_ir.json"
        ir_json_path.write_text(json.dumps(ir_payload, indent=2), encoding="utf-8")
        excel_path = bundle_dir / f"{sequence_name}_from_awl.xlsx"
        _export_excel_from_ir(ir_payload, excel_path)

        result = export_conversion_bundle_from_ir(
            sequence_name=sequence_name,
            ir_payload=ir_payload,
            source_name=source.name,
            output_dir=str(bundle_dir_relative),
        )
        # Preserve project-level warnings/metadata (e.g. missing called blocks) which would
        # otherwise be lost when re-analyzing from IR.
        analysis_json_path = bundle_dir / f"{sequence_name}_analysis.json"
        analysis_json_path.write_text(json.dumps(analysis, indent=2), encoding="utf-8")
        generated += 1
        print(f"[OK] {source.name} -> {result['outputDirectory']}")
        print(f"[IR] {ir_json_path}")
        print(f"[EXCEL] {excel_path}")

    print(f"Done. Generated bundles: {generated}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
